from __future__ import annotations
import csv, hashlib, html, io, os, secrets, shutil, smtplib, uuid
from datetime import datetime, timedelta
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo
from urllib.parse import quote, urlsplit
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session
from app import models
from app.database import SessionLocal, engine, get_db
from app.dependencies import get_usuario_actual, hashear_password
from app.permissions import allowed_set_ids, can_access_set, is_admin, require_access_set, require_admin_set, require_finance_set, require_staff_set

router = APIRouter(prefix="/v2", tags=["Funciones ampliadas"])
TZ = ZoneInfo("America/Bogota")
MAX_UPLOAD = 10 * 1024 * 1024


def localnow() -> datetime:
    return datetime.now(TZ).replace(tzinfo=None)


def dt(value: str | datetime) -> datetime:
    value = datetime.fromisoformat(value.replace("Z", "+00:00")) if isinstance(value, str) else value
    return value.astimezone(TZ).replace(tzinfo=None) if value.tzinfo else value


def notify(db: Session, uid: int, title: str, message: str) -> None:
    db.add(models.Notificacion(usuario_id=uid, titulo=title, mensaje=message))
    db.flush()
    from app.services.notifications import deliver_external
    deliver_external(db, uid, title, message)


def units_for_user(db: Session, user: models.Usuario) -> list[models.Inmueble]:
    return db.query(models.Inmueble).filter((models.Inmueble.propietario_id == user.id) | (models.Inmueble.inquilino_id == user.id)).all()


def unit_access(db: Session, user: models.Usuario, unit_id: int) -> models.Inmueble:
    unit = db.query(models.Inmueble).filter_by(id=unit_id).first()
    if not unit:
        raise HTTPException(404, "Inmueble no encontrado")
    if not is_admin(user) and user.id not in (unit.propietario_id, unit.inquilino_id):
        raise HTTPException(403, "No tienes acceso a este inmueble")
    if is_admin(user):
        require_admin_set(db, user, unit.conjunto_id)
    return unit


def admin_only(user: models.Usuario) -> None:
    if not is_admin(user):
        raise HTTPException(403, "Solo administración puede realizar esta acción")


def model_data(row):
    return {column.name: getattr(row, column.name) for column in row.__table__.columns}


def workbook_response(wb: Workbook, filename: str):
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def safe_cell(value: Any):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def xlsx_template(kind: str) -> Workbook:
    headers = {
        "residents": ["nombre_completo", "cedula", "email", "telefono", "password", "inmueble_torre", "inmueble_numero", "estado_inmueble"],
        "properties": ["torre", "numero", "propietario_email", "inquilino_email", "estado"],
        "payments": ["inmueble_id", "concepto", "periodo", "monto", "fecha_vencimiento", "estado", "referencia"],
    }
    if kind not in headers:
        raise HTTPException(404, "Plantilla no encontrada")
    wb = Workbook(); ws = wb.active; ws.title = kind.capitalize(); ws.append(headers[kind]); ws.freeze_panes = "A2"
    return wb


@router.get("/summary")
def summary(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if is_admin(user):
        admin_only(user); ids = allowed_set_ids(db, user)
        unit_ids = [x.id for x in db.query(models.Inmueble.id).filter(models.Inmueble.conjunto_id.in_(ids)).all()]
    else:
        unit_ids = [x.id for x in units_for_user(db, user)]
        ids = {x.conjunto_id for x in units_for_user(db, user)}
    return {
        "conjuntos": len(ids), "inmuebles": len(unit_ids),
        "reservas_pendientes": db.query(models.Reserva).join(models.ZonaComun).filter(models.ZonaComun.conjunto_id.in_(ids), models.Reserva.estado == "pendiente").count() if is_admin(user) and ids else 0,
        "pqrs_pendientes": db.query(models.PQRS).join(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids), models.PQRS.estado != "resuelto").count() if is_admin(user) and ids else 0,
        "cuotas_pendientes": db.query(models.Cuota).filter(models.Cuota.inmueble_id.in_(unit_ids), models.Cuota.estado == "pendiente").count() if unit_ids else 0,
        "saldo_pendiente": str(sum((Decimal(str(x.monto)) for x in db.query(models.Cuota).filter(models.Cuota.inmueble_id.in_(unit_ids), models.Cuota.estado == "pendiente").all()), Decimal("0"))),
        "pagos_por_verificar": db.query(models.Pago).join(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids), models.Pago.estado == "pendiente").count() if is_admin(user) and ids else 0,
        "incidentes_abiertos": db.query(models.Incidente).filter(models.Incidente.conjunto_id.in_(ids), models.Incidente.estado.in_(["abierto", "en_proceso"])).count() if is_admin(user) and ids else db.query(models.Incidente).filter(models.Incidente.solicitante_id == user.id, models.Incidente.estado.in_(["abierto", "en_proceso"])).count(),
        "paquetes_pendientes": db.query(models.Paquete).filter(models.Paquete.inmueble_id.in_(unit_ids), models.Paquete.retirado.is_(False)).count() if unit_ids else 0,
    }


@router.get("/memberships")
def memberships(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if user.es_super_admin:
        rows = db.query(models.ConjuntoMembership).order_by(models.ConjuntoMembership.conjunto_id).all()
    elif user.rol in ("admin", "portero"):
        rows = db.query(models.ConjuntoMembership).filter_by(usuario_id=user.id).all()
    else:
        raise HTTPException(403, "Sin permiso")
    return [{"id": r.id, "conjunto_id": r.conjunto_id, "usuario_id": r.usuario_id, "rol": r.rol,
             "nombre": db.query(models.Usuario.nombre_completo).filter_by(id=r.usuario_id).scalar(),
             "email": db.query(models.Usuario.email).filter_by(id=r.usuario_id).scalar(),
             "conjunto": db.query(models.Conjunto.nombre).filter_by(id=r.conjunto_id).scalar()} for r in rows]


@router.post("/memberships", status_code=201)
def add_membership(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if not user.es_super_admin: raise HTTPException(403, "Solo superadministración puede asignar equipos por conjunto")
    uid, cid, role = int(data.get("usuario_id", 0)), int(data.get("conjunto_id", 0)), data.get("rol")
    target = db.query(models.Usuario).filter_by(id=uid).first(); group = db.query(models.Conjunto).filter_by(id=cid).first()
    if not target or not group: raise HTTPException(404, "Usuario o conjunto no encontrado")
    if role not in ("admin", "portero", "contador"): raise HTTPException(422, "El rol de conjunto debe ser admin, portero o contador")
    conflict = db.query(models.ConjuntoMembership).filter(models.ConjuntoMembership.usuario_id == uid, models.ConjuntoMembership.rol != role).first()
    if conflict: raise HTTPException(409, "Cada cuenta operativa conserva el mismo rol en todos sus conjuntos; elimina la asignación anterior para cambiarlo")
    row = db.query(models.ConjuntoMembership).filter_by(usuario_id=uid, conjunto_id=cid).first()
    if row: row.rol = role
    else: row = models.ConjuntoMembership(usuario_id=uid, conjunto_id=cid, rol=role); db.add(row)
    target.rol = role; db.commit(); db.refresh(row)
    return {"id": row.id, "usuario_id": uid, "conjunto_id": cid, "rol": role}


@router.delete("/memberships/{membership_id}", status_code=204)
def delete_membership(membership_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if not user.es_super_admin: raise HTTPException(403, "Solo superadministración puede quitar asignaciones")
    row = db.query(models.ConjuntoMembership).filter_by(id=membership_id).first()
    if not row: raise HTTPException(404, "Asignación no encontrada")
    target = db.query(models.Usuario).filter_by(id=row.usuario_id).first()
    role = row.rol
    db.delete(row); db.flush()
    remaining = db.query(models.ConjuntoMembership).filter_by(usuario_id=row.usuario_id, rol=role).first()
    if target and not target.es_super_admin and not remaining and target.rol == role:
        target.rol = "residente"
    db.commit()
    from fastapi import Response
    return Response(status_code=204)


@router.get("/payments")
def list_payments(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    q = db.query(models.Pago)
    if is_admin(user) or user.rol == "contador":
        ids = allowed_set_ids(db, user)
        if not ids: return []
        q = q.join(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids))
    else:
        q = q.filter(models.Pago.inmueble_id.in_([x.id for x in units_for_user(db, user)]))
    return [model_data(x) for x in q.order_by(models.Pago.creado_en.desc()).all()]


@router.post("/payments", status_code=201)
def request_payment(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    cuota = db.query(models.Cuota).filter_by(id=int(data.get("cuota_id", 0))).first()
    if not cuota: raise HTTPException(404, "Cuota no encontrada")
    unit = unit_access(db, user, cuota.inmueble_id)
    if is_admin(user): raise HTTPException(400, "Use el registro manual de administración")
    if cuota.estado == "pagada": raise HTTPException(409, "La cuota ya figura pagada")
    if db.query(models.Pago).filter_by(cuota_id=cuota.id, estado="pendiente").first(): raise HTTPException(409, "Ya hay un pago pendiente de verificación")
    method = str(data.get("metodo", "transferencia"))
    if method not in ("transferencia", "efectivo", "consignacion", "otro"): raise HTTPException(422, "Método no permitido")
    row = models.Pago(cuota_id=cuota.id, inmueble_id=unit.id, solicitado_por=user.id, monto=cuota.monto, metodo=method, referencia=str(data.get("referencia", ""))[:120], estado="pendiente")
    db.add(row)
    admins = db.query(models.Usuario).filter(models.Usuario.es_super_admin.is_(True)).all()
    admins += db.query(models.Usuario).join(models.ConjuntoMembership, models.ConjuntoMembership.usuario_id == models.Usuario.id).filter(models.ConjuntoMembership.conjunto_id == unit.conjunto_id, models.ConjuntoMembership.rol.in_(("admin", "contador"))).all()
    for a in {x.id: x for x in admins}.values(): notify(db, a.id, "Pago por verificar", f"Pago reportado de ${cuota.monto} para cuota #{cuota.id}")
    db.commit(); db.refresh(row); return model_data(row)


@router.post("/payments/manual", status_code=201)
def manual_payment(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    cuota = db.query(models.Cuota).filter_by(id=int(data.get("cuota_id", 0))).first()
    if not cuota: raise HTTPException(404, "Cuota no encontrada")
    conjunto_id = db.query(models.Inmueble.conjunto_id).filter_by(id=cuota.inmueble_id).scalar()
    if user.rol == "contador": require_finance_set(db, user, conjunto_id)
    else: require_admin_set(db, user, conjunto_id)
    if cuota.estado == "pagada": raise HTTPException(409, "La cuota ya está pagada")
    row = models.Pago(cuota_id=cuota.id, inmueble_id=cuota.inmueble_id, solicitado_por=user.id, verificado_por=user.id, monto=cuota.monto, metodo=str(data.get("metodo", "efectivo")), referencia=str(data.get("referencia", ""))[:120], estado="verificado", notas=str(data.get("notas", ""))[:1000], verificado_en=localnow())
    cuota.estado = "pagada"; cuota.pagada_en = row.verificado_en; db.add(row)
    for uid in {db.query(models.Inmueble.propietario_id).filter_by(id=cuota.inmueble_id).scalar(), db.query(models.Inmueble.inquilino_id).filter_by(id=cuota.inmueble_id).scalar()} - {None}: notify(db, uid, "Pago registrado", f"Se registró el pago de ${cuota.monto} para {cuota.concepto}")
    db.commit(); db.refresh(row); return model_data(row)


@router.patch("/payments/{payment_id}")
def verify_payment(payment_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Pago).filter_by(id=payment_id).first()
    if not row: raise HTTPException(404, "Pago no encontrado")
    unit = db.query(models.Inmueble).filter_by(id=row.inmueble_id).first()
    if user.rol == "contador": require_finance_set(db, user, unit.conjunto_id)
    else: require_admin_set(db, user, unit.conjunto_id)
    if row.estado != "pendiente": raise HTTPException(409, "El pago ya fue conciliado")
    state = data.get("estado")
    if state not in ("verificado", "rechazado"): raise HTTPException(422, "Estado de conciliación inválido")
    row.estado = state; row.verificado_por = user.id; row.verificado_en = localnow(); row.notas = str(data.get("notas", ""))[:1000]
    if state == "verificado":
        fee = db.query(models.Cuota).filter_by(id=row.cuota_id).first(); fee.estado = "pagada"; fee.pagada_en = row.verificado_en
    notify(db, row.solicitado_por, "Pago " + state, f"El pago #{row.id} fue {state}. {row.notas or ''}")
    db.commit(); db.refresh(row); return model_data(row)


@router.get("/payments/{payment_id}/receipt", response_class=HTMLResponse)
def payment_receipt(payment_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Pago).filter_by(id=payment_id).first()
    if not row: raise HTTPException(404, "Recibo no encontrado")
    unit = unit_access(db, user, row.inmueble_id)
    if row.estado != "verificado": raise HTTPException(409, "El recibo se habilita cuando la administración verifica el pago")
    fee = db.query(models.Cuota).filter_by(id=row.cuota_id).first(); group = db.query(models.Conjunto).filter_by(id=unit.conjunto_id).first()
    content = f'''<!doctype html><html lang="es"><meta charset="utf-8"><title>Recibo IRIS #{row.id}</title><style>body{{font:16px Arial;max-width:720px;margin:50px auto;color:#183d43}}header,footer{{border-bottom:2px solid #176d69;padding:14px}}table{{width:100%;border-collapse:collapse}}td{{padding:12px;border-bottom:1px solid #ddd}}button{{padding:10px 18px;background:#176d69;color:#fff;border:0;cursor:pointer}}@media print{{button{{display:none}}}}</style><header><h1>Recibo de pago</h1><b>{html.escape(group.nombre if group else 'IRIS')}</b></header><p>Recibo #{row.id} · {row.verificado_en.strftime('%Y-%m-%d %H:%M') if row.verificado_en else ''}</p><table><tr><td>Inmueble</td><td>{html.escape((unit.torre or '')+' '+unit.numero)}</td></tr><tr><td>Concepto</td><td>{html.escape(fee.concepto)}</td></tr><tr><td>Periodo</td><td>{html.escape(fee.periodo)}</td></tr><tr><td>Monto</td><td>${Decimal(str(row.monto)):,.2f} COP</td></tr><tr><td>Método</td><td>{html.escape(row.metodo)}</td></tr><tr><td>Referencia</td><td>{html.escape(row.referencia or '—')}</td></tr><tr><td>Estado</td><td>VERIFICADO</td></tr></table><footer>Comprobante generado por IRIS. <button onclick="window.print()">Imprimir / Guardar PDF</button></footer></html>'''
    return HTMLResponse(content)


@router.get("/incidents")
def list_incidents(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    q = db.query(models.Incidente)
    if is_admin(user):
        ids = allowed_set_ids(db, user)
        if not ids: return []
        q = q.filter(models.Incidente.conjunto_id.in_(ids))
    elif user.rol == "portero":
        ids = allowed_set_ids(db, user); q = q.filter(models.Incidente.conjunto_id.in_(ids))
    else: q = q.filter(models.Incidente.solicitante_id == user.id)
    return [model_data(x) for x in q.order_by(models.Incidente.creado_en.desc()).limit(500).all()]


@router.post("/incidents", status_code=201)
def create_incident(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    title, desc = str(data.get("titulo", "")).strip(), str(data.get("descripcion", "")).strip()
    if not title or not desc: raise HTTPException(422, "Título y descripción son obligatorios")
    if is_admin(user):
        cid = int(data.get("conjunto_id", 0)); require_admin_set(db, user, cid); unit_id = data.get("inmueble_id")
    else:
        unit = unit_access(db, user, int(data.get("inmueble_id", 0))); cid = unit.conjunto_id; unit_id = unit.id
    row = models.Incidente(conjunto_id=cid, inmueble_id=unit_id, solicitante_id=user.id, categoria=str(data.get("categoria", "mantenimiento"))[:50], titulo=title[:160], descripcion=desc[:5000], prioridad="media", estado="abierto")
    db.add(row); db.commit(); db.refresh(row)
    admins = db.query(models.Usuario).filter(models.Usuario.es_super_admin.is_(True)).all() + db.query(models.Usuario).join(models.ConjuntoMembership, models.ConjuntoMembership.usuario_id == models.Usuario.id).filter(models.ConjuntoMembership.conjunto_id == cid, models.ConjuntoMembership.rol == "admin").all()
    for a in {x.id: x for x in admins}.values(): notify(db, a.id, "Nuevo incidente", title)
    db.commit(); db.refresh(row); return model_data(row)


@router.patch("/incidents/{incident_id}")
def update_incident(incident_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Incidente).filter_by(id=incident_id).first()
    if not row: raise HTTPException(404, "Incidente no encontrado")
    require_admin_set(db, user, row.conjunto_id)
    if data.get("estado") is not None:
        if data["estado"] not in ("abierto", "en_proceso", "resuelto", "cerrado"): raise HTTPException(422, "Estado inválido")
        row.estado = data["estado"]
    if data.get("prioridad") is not None:
        if data["prioridad"] not in ("baja", "media", "alta", "urgente"): raise HTTPException(422, "Prioridad inválida")
        row.prioridad = data["prioridad"]
    if data.get("asignado_a") is not None:
        assigned = db.query(models.Usuario).filter_by(id=int(data["asignado_a"])).first()
        if not assigned: raise HTTPException(404, "Usuario asignado no encontrado")
        row.asignado_a = assigned.id
    row.actualizado_en = localnow(); notify(db, row.solicitante_id, "Actualización de incidente", f"{row.titulo}: {row.estado}")
    db.commit(); db.refresh(row); return model_data(row)


async def save_image(file: UploadFile) -> str:
    if not file.content_type or file.content_type.lower() not in ("image/jpeg", "image/png", "image/webp"):
        raise HTTPException(415, "Usa una imagen JPG, PNG o WebP")
    content = await file.read(MAX_UPLOAD + 1)
    if len(content) > MAX_UPLOAD: raise HTTPException(413, "El archivo excede 10 MB")
    signatures = {"image/jpeg": content.startswith(b"\xff\xd8\xff"), "image/png": content.startswith(b"\x89PNG\r\n\x1a\n"), "image/webp": len(content) > 12 and content.startswith(b"RIFF") and content[8:12] == b"WEBP"}
    if not signatures.get(file.content_type.lower()): raise HTTPException(415, "El contenido del archivo no coincide con su tipo")
    ext = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}[file.content_type.lower()]
    root = Path(os.getenv("UPLOAD_DIR", "./uploads")); root.mkdir(parents=True, exist_ok=True)
    name = uuid.uuid4().hex + ext; (root / name).write_bytes(content)
    return "/v2/files/" + name


@router.post("/incidents/{incident_id}/photo")
async def incident_photo(incident_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Incidente).filter_by(id=incident_id).first()
    if not row: raise HTTPException(404, "Incidente no encontrado")
    if is_admin(user): require_admin_set(db, user, row.conjunto_id)
    elif row.solicitante_id != user.id: raise HTTPException(403, "Sin permiso")
    row.foto_url = await save_image(file); db.commit(); return {"foto_url": row.foto_url}


@router.get("/emergency-contacts")
def list_contacts(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    ids = {i.conjunto_id for i in units_for_user(db, user)}
    if is_admin(user) or user.rol == "portero": ids = allowed_set_ids(db, user)
    q = db.query(models.ContactoEmergencia)
    if is_admin(user): q = q.filter(models.ContactoEmergencia.conjunto_id.in_(ids))
    elif user.rol == "portero": q = q.filter(models.ContactoEmergencia.conjunto_id.in_(ids), models.ContactoEmergencia.visible.is_(True))
    else: q = q.filter(models.ContactoEmergencia.conjunto_id.in_(ids), (models.ContactoEmergencia.usuario_id == user.id) | (models.ContactoEmergencia.visible.is_(True)))
    return [model_data(x) for x in q.order_by(models.ContactoEmergencia.conjunto_id, models.ContactoEmergencia.inmueble_id).all()]


@router.post("/emergency-contacts", status_code=201)
def create_contact(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    unit = unit_access(db, user, int(data.get("inmueble_id", 0)))
    name, phone = str(data.get("nombre", "")).strip(), str(data.get("telefono", "")).strip()
    if not name or not phone: raise HTTPException(422, "Nombre y teléfono son obligatorios")
    row = models.ContactoEmergencia(conjunto_id=unit.conjunto_id, inmueble_id=unit.id, usuario_id=user.id, nombre=name[:120], telefono=phone[:30], parentesco=str(data.get("parentesco", ""))[:60], visible=bool(data.get("visible", False)))
    db.add(row); db.commit(); db.refresh(row); return model_data(row)


@router.delete("/emergency-contacts/{contact_id}", status_code=204)
def delete_contact(contact_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.ContactoEmergencia).filter_by(id=contact_id).first()
    if not row: raise HTTPException(404, "Contacto no encontrado")
    if not is_admin(user) and row.usuario_id != user.id: raise HTTPException(403, "Sin permiso")
    if is_admin(user): require_admin_set(db, user, row.conjunto_id)
    db.delete(row); db.commit()
    from fastapi import Response
    return Response(status_code=204)


@router.post("/announcements/{announcement_id}/read")
def mark_announcement_read(announcement_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Comunicado).filter_by(id=announcement_id, publicado=True).first()
    if not row: raise HTTPException(404, "Comunicado no encontrado")
    require_access_set(db, user, row.conjunto_id)
    receipt = db.query(models.ComunicadoLectura).filter_by(comunicado_id=row.id, usuario_id=user.id).first()
    if not receipt: db.add(models.ComunicadoLectura(comunicado_id=row.id, usuario_id=user.id)); db.commit()
    return {"leido": True}


@router.get("/announcements/{announcement_id}/reads")
def announcement_reads(announcement_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Comunicado).filter_by(id=announcement_id).first()
    if not row: raise HTTPException(404, "Comunicado no encontrado")
    require_admin_set(db, user, row.conjunto_id)
    return {"leidos": db.query(models.ComunicadoLectura).filter_by(comunicado_id=row.id).count()}


@router.get("/my-units")
def my_units(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if is_admin(user) or user.rol == "portero": raise HTTPException(403, "Este endpoint es para residentes")
    return [{"id": u.id, "conjunto_id": u.conjunto_id, "conjunto": db.query(models.Conjunto.nombre).filter_by(id=u.conjunto_id).scalar(), "torre": u.torre, "numero": u.numero, "etiqueta": (f"Torre {u.torre} · " if u.torre else "") + f"Apto {u.numero}", "rol": "propietario" if u.propietario_id == user.id else "inquilino"} for u in units_for_user(db, user)]


@router.get("/operational-units")
def operational_units(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if not (is_admin(user) or user.rol == "portero"):
        raise HTTPException(403, "Este catálogo es para administración y portería")
    ids = allowed_set_ids(db, user)
    rows = db.query(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids)).order_by(models.Inmueble.torre, models.Inmueble.numero).all() if ids else []
    return [{"id": row.id, "conjunto_id": row.conjunto_id, "conjunto": db.query(models.Conjunto.nombre).filter_by(id=row.conjunto_id).scalar(), "torre": row.torre, "numero": row.numero, "etiqueta": (f"Torre {row.torre} · " if row.torre else "") + f"Apto {row.numero}"} for row in rows]


@router.get("/reservation-blocks")
def list_blocks(zona_id: int | None = None, desde: datetime | None = None, hasta: datetime | None = None, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if (desde is None) != (hasta is None): raise HTTPException(422, "Indica ambas fechas del filtro")
    q = db.query(models.BloqueoReserva).join(models.ZonaComun)
    if zona_id is not None:
        zone = db.query(models.ZonaComun).filter_by(id=zona_id).first()
        if not zone: raise HTTPException(404, "Zona no encontrada")
        require_access_set(db, user, zone.conjunto_id)
        q = q.filter(models.BloqueoReserva.zona_id == zona_id)
    elif is_admin(user) or user.rol == "portero":
        ids = allowed_set_ids(db, user); q = q.filter(models.ZonaComun.conjunto_id.in_(ids))
    else:
        ids = {x.conjunto_id for x in units_for_user(db, user)}; q = q.filter(models.ZonaComun.conjunto_id.in_(ids))
    if desde is not None and hasta is not None:
        start, end = dt(desde), dt(hasta)
        if end <= start: raise HTTPException(422, "Rango inválido")
        q = q.filter(models.BloqueoReserva.fecha_inicio < end, models.BloqueoReserva.fecha_fin > start)
    return [model_data(x) for x in q.order_by(models.BloqueoReserva.fecha_inicio).limit(500).all()]


@router.post("/reservation-blocks", status_code=201)
def create_block(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    zone = db.query(models.ZonaComun).filter_by(id=int(data.get("zona_id", 0))).first()
    if not zone: raise HTTPException(404, "Zona no encontrada")
    require_admin_set(db, user, zone.conjunto_id)
    start, end = dt(data["fecha_inicio"]), dt(data["fecha_fin"])
    if end <= start: raise HTTPException(422, "La fecha de fin debe ser posterior")
    reason = str(data.get("motivo", "Mantenimiento")).strip()
    if not reason: raise HTTPException(422, "Indica el motivo del bloqueo")
    row = models.BloqueoReserva(zona_id=zone.id, fecha_inicio=start, fecha_fin=end, motivo=reason[:255], creado_por=user.id)
    db.add(row); db.commit(); db.refresh(row); return model_data(row)


@router.delete("/reservation-blocks/{block_id}", status_code=204)
def delete_block(block_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.BloqueoReserva).filter_by(id=block_id).first()
    if not row: raise HTTPException(404, "Bloqueo no encontrado")
    zone = db.query(models.ZonaComun).filter_by(id=row.zona_id).first(); require_admin_set(db, user, zone.conjunto_id)
    db.delete(row); db.commit()
    from fastapi import Response
    return Response(status_code=204)


@router.get("/audit")
def audit(limit: int = 100, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    admin_only(user); ids = allowed_set_ids(db, user)
    if not ids and not user.es_super_admin: return []
    q = db.query(models.Auditoria)
    if not user.es_super_admin: q = q.filter(models.Auditoria.conjunto_id.in_(ids))
    return [model_data(x) for x in q.order_by(models.Auditoria.id.desc()).limit(max(1, min(limit, 500))).all()]


@router.get("/export/{kind}.xlsx")
def export_xlsx(kind: str, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    admin_only(user); ids = allowed_set_ids(db, user)
    if kind == "payments":
        rows = db.query(models.Pago).join(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids)).order_by(models.Pago.creado_en.desc()).all() if ids else []
        headers = ["id", "cuota_id", "inmueble_id", "monto", "metodo", "referencia", "estado", "solicitado_por", "verificado_en"]
    elif kind == "residents":
        rows = db.query(models.Usuario).join(models.Inmueble, (models.Inmueble.propietario_id == models.Usuario.id) | (models.Inmueble.inquilino_id == models.Usuario.id)).filter(models.Inmueble.conjunto_id.in_(ids)).distinct().all() if ids else []
        headers = ["id", "nombre_completo", "cedula", "email", "telefono", "rol"]
    elif kind == "properties":
        rows = db.query(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids)).all() if ids else []
        headers = ["id", "conjunto_id", "torre", "numero", "propietario_id", "inquilino_id", "estado"]
    elif kind == "incidents":
        rows = db.query(models.Incidente).filter(models.Incidente.conjunto_id.in_(ids)).all() if ids else []
        headers = ["id", "conjunto_id", "inmueble_id", "solicitante_id", "titulo", "prioridad", "estado", "creado_en"]
    else: raise HTTPException(404, "Tipo de reporte no soportado")
    wb = Workbook(); ws = wb.active; ws.title = kind[:31]; ws.append(headers)
    for row in rows: ws.append([safe_cell(getattr(row, key, None)) for key in headers])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return workbook_response(wb, f"iris-{kind}.xlsx")


@router.get("/import/template/{kind}")
def import_template(kind: str, user=Depends(get_usuario_actual)):
    admin_only(user); return workbook_response(xlsx_template(kind), f"plantilla-{kind}.xlsx")


@router.post("/import/{kind}")
async def import_xlsx(kind: str, file: UploadFile = File(...), conjunto_id: int | None = None, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    admin_only(user)
    if kind not in ("residents", "properties", "payments"): raise HTTPException(404, "Importación no soportada")
    if conjunto_id is None: raise HTTPException(422, "Se requiere conjunto_id")
    require_admin_set(db, user, conjunto_id)
    raw = await file.read(MAX_UPLOAD + 1)
    if len(raw) > MAX_UPLOAD: raise HTTPException(413, "Archivo excede 10 MB")
    if not (file.filename or "").lower().endswith(".xlsx"): raise HTTPException(415, "Sube un archivo .xlsx")
    try:
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        iterator = ws.iter_rows(values_only=True); headers = [str(x or "").strip().lower() for x in next(iterator)]
    except Exception: raise HTTPException(400, "No se pudo leer el XLSX")
    result = {"creados": 0, "actualizados": 0, "errores": []}
    for line, values in enumerate(iterator, start=2):
        data = dict(zip(headers, values))
        if not any(v is not None for v in values): continue
        try:
            if kind == "residents":
                email = str(data.get("email", "")).strip().lower(); doc = str(data.get("cedula", "")).strip()
                name = str(data.get("nombre_completo", "")).strip(); password = str(data.get("password", ""))
                torre = str(data.get("inmueble_torre", "")).strip() or None; number = str(data.get("inmueble_numero", "")).strip()
                if not email or not doc or not name or len(password) < 10 or not number: raise ValueError("nombre, cédula, email, password (10+ caracteres) e inmueble_numero son obligatorios")
                group_id = int(data.get("conjunto_id") or conjunto_id or 0); require_admin_set(db, user, group_id)
                existing = db.query(models.Usuario).filter(models.Usuario.email == email).first()
                if existing: raise ValueError("email ya registrado; no se sobrescribe usuario existente")
                if db.query(models.Usuario).filter(models.Usuario.cedula == doc).first(): raise ValueError("cédula ya registrada")
                resident = models.Usuario(nombre_completo=name[:100], cedula=doc[:20], email=email, telefono=str(data.get("telefono", ""))[:20] or None, password_hash=hashear_password(password), rol="residente", email_verificado=os.getenv("APP_ENV", "development") != "production", cambio_password_requerido=True)
                db.add(resident); db.flush()
                state = str(data.get("estado_inmueble") or "ocupado").lower()
                if state not in ("ocupado", "desocupado", "en_venta"): raise ValueError("estado_inmueble debe ser ocupado, desocupado o en_venta")
                if db.query(models.Inmueble).filter_by(conjunto_id=group_id, torre=torre, numero=number).first(): raise ValueError("el inmueble ya existe")
                db.add(models.Inmueble(conjunto_id=group_id, torre=torre, numero=number, propietario_id=resident.id, estado=state))
                if os.getenv("APP_ENV", "development").lower() == "production":
                    from app.security_tokens import issue_token
                    from app.services.notifications import deliver_external
                    token = issue_token(db, resident, "verify", 24 * 60)
                    link = f"{os.getenv('API_PUBLIC_URL', 'http://localhost:8000')}/auth/verify-email?token={quote(token)}"
                    deliver_external(db, resident.id, "Activa tu cuenta IRIS", f"Abre este enlace para activar tu cuenta: {link}", allow_unverified=True)
                result["creados"] += 1
            elif kind == "properties":
                owners = db.query(models.Usuario).filter_by(email=str(data.get("propietario_email", "")).strip().lower()).first()
                if not owners: raise ValueError("propietario_email no existe")
                tenant_email = str(data.get("inquilino_email", "")).strip().lower(); tenant = db.query(models.Usuario).filter_by(email=tenant_email).first() if tenant_email else None
                if tenant_email and not tenant: raise ValueError("inquilino_email no existe")
                number = str(data.get("numero", "")).strip()
                if not number: raise ValueError("numero es obligatorio")
                tower = str(data.get("torre", "")).strip() or None
                if db.query(models.Inmueble).filter_by(conjunto_id=conjunto_id, torre=tower, numero=number).first(): raise ValueError("el inmueble ya existe")
                state = str(data.get("estado") or "ocupado").lower()
                if state not in ("ocupado", "desocupado", "en_venta"): raise ValueError("estado debe ser ocupado, desocupado o en_venta")
                db.add(models.Inmueble(conjunto_id=conjunto_id, torre=tower, numero=number, propietario_id=owners.id, inquilino_id=tenant.id if tenant else None, estado=state))
                result["creados"] += 1
            else:
                unit = db.query(models.Inmueble).filter_by(id=int(data.get("inmueble_id", 0)), conjunto_id=conjunto_id).first()
                if not unit: raise ValueError("inmueble_id no pertenece al conjunto")
                due = datetime.strptime(str(data.get("fecha_vencimiento")), "%Y-%m-%d").date(); amount = Decimal(str(data.get("monto")))
                if amount <= 0: raise ValueError("monto debe ser mayor que cero")
                state = str(data.get("estado") or "pendiente").lower()
                if state not in ("pendiente", "pagada"): raise ValueError("estado debe ser pendiente o pagada")
                concept = str(data.get("concepto", "Administración"))[:160]; period = str(data.get("periodo", ""))[:20]
                if not period: raise ValueError("periodo es obligatorio")
                if db.query(models.Cuota).filter_by(inmueble_id=unit.id, concepto=concept, periodo=period).first(): raise ValueError("ya existe una cuota con ese concepto y periodo")
                fee = models.Cuota(inmueble_id=unit.id, concepto=concept, periodo=period, monto=amount, fecha_vencimiento=due, estado=state, creado_por=user.id, pagada_en=localnow() if state == "pagada" else None)
                db.add(fee); db.flush()
                if state == "pagada":
                    db.add(models.Pago(cuota_id=fee.id, inmueble_id=unit.id, solicitado_por=user.id, verificado_por=user.id, monto=amount, metodo="otro", referencia=str(data.get("referencia", ""))[:120] or None, estado="verificado", notas="Importado desde XLSX", verificado_en=fee.pagada_en))
                result["creados"] += 1
            db.commit()
        except Exception as exc:
            db.rollback(); result["errores"].append({"fila": line, "detalle": str(exc)[:200]})
    return result


@router.get("/admin/backup")
def sqlite_backup(user=Depends(get_usuario_actual)):
    admin_only(user)
    if not engine.url.drivername.startswith("sqlite"): raise HTTPException(501, "Para PostgreSQL/MySQL configura snapshots gestionados del proveedor de base de datos")
    source = engine.raw_connection(); output = io.BytesIO()
    try:
        temp_path = Path("/tmp") / f"iris-backup-{uuid.uuid4().hex}.db"; destination = __import__("sqlite3").connect(str(temp_path)); source.driver_connection.backup(destination); destination.close()
        output.write(temp_path.read_bytes()); temp_path.unlink(missing_ok=True); output.seek(0)
    finally: source.close()
    return StreamingResponse(output, media_type="application/vnd.sqlite3", headers={"Content-Disposition": f'attachment; filename="iris-backup-{localnow():%Y%m%d-%H%M}.db"', "Cache-Control": "no-store"})


@router.get("/push/public-key")
def push_public_key(user=Depends(get_usuario_actual)):
    return {"publicKey": os.getenv("VAPID_PUBLIC_KEY", "")}


@router.post("/push/subscriptions", status_code=201)
def push_subscribe(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    endpoint = str(data.get("endpoint", "")); keys = data.get("keys") or {}
    if len(endpoint) > 2048 or len(str(keys.get("p256dh", ""))) > 300 or len(str(keys.get("auth", ""))) > 100:
        raise HTTPException(422, "Suscripción Web Push demasiado grande")
    try:
        parsed = urlsplit(endpoint); port = parsed.port
    except ValueError:
        raise HTTPException(422, "URL Web Push inválida")
    hostname = (parsed.hostname or "").lower()
    allowed_push_domains = ("googleapis.com", "mozilla.com", "push.apple.com", "notify.windows.com")
    trusted_push_host = any(hostname == domain or hostname.endswith("." + domain) for domain in allowed_push_domains)
    if parsed.scheme != "https" or parsed.username or parsed.password or port not in (None, 443) or not trusted_push_host or not keys.get("p256dh") or not keys.get("auth"):
        raise HTTPException(422, "Suscripción Web Push inválida o proveedor no admitido")
    row = db.query(models.PushSubscription).filter_by(endpoint=endpoint).first()
    if row: row.usuario_id=user.id; row.p256dh=keys["p256dh"]; row.auth=keys["auth"]
    else: db.add(models.PushSubscription(usuario_id=user.id, endpoint=endpoint, p256dh=keys["p256dh"], auth=keys["auth"]))
    db.commit(); return {"ok": True}


@router.delete("/push/subscriptions", status_code=204)
def push_unsubscribe(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    endpoint = str(data.get("endpoint", "")); db.query(models.PushSubscription).filter_by(usuario_id=user.id, endpoint=endpoint).delete(); db.commit()
    from fastapi import Response
    return Response(status_code=204)


@router.post("/packages/{package_id}/photo")
async def package_photo(package_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row=db.query(models.Paquete).filter_by(id=package_id).first()
    if not row: raise HTTPException(404, "Paquete no encontrado")
    unit=db.query(models.Inmueble).filter_by(id=row.inmueble_id).first(); require_staff_set(db,user,unit.conjunto_id)
    row.foto_url=await save_image(file); db.commit(); return {"foto_url":row.foto_url}


@router.post("/announcements/{announcement_id}/attachment")
async def announcement_attachment(announcement_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Comunicado).filter_by(id=announcement_id).first()
    if not row: raise HTTPException(404, "Comunicado no encontrado")
    require_admin_set(db, user, row.conjunto_id)
    mime = (file.content_type or "").lower()
    content = await file.read(MAX_UPLOAD + 1)
    if len(content) > MAX_UPLOAD: raise HTTPException(413, "Archivo excede 10 MB")
    allowed = {"application/pdf": (b"%PDF", ".pdf"), "image/jpeg": (b"\xff\xd8\xff", ".jpg"), "image/png": (b"\x89PNG\r\n\x1a\n", ".png"), "image/webp": (b"RIFF", ".webp")}
    if mime not in allowed or not content.startswith(allowed[mime][0]) or (mime == "image/webp" and content[8:12] != b"WEBP"):
        raise HTTPException(415, "Adjunta PDF, JPG, PNG o WebP válido")
    root = Path(os.getenv("UPLOAD_DIR", "./uploads")); root.mkdir(parents=True, exist_ok=True)
    name = uuid.uuid4().hex + allowed[mime][1]; (root / name).write_bytes(content)
    row.adjunto_url = "/v2/files/" + name; db.commit()
    return {"adjunto_url": row.adjunto_url}


@router.get("/files/{filename}")
def authenticated_file(filename: str, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if Path(filename).name != filename or filename.startswith("."):
        raise HTTPException(404, "Archivo no encontrado")
    url = "/v2/files/" + filename
    incident = db.query(models.Incidente).filter_by(foto_url=url).first()
    package = db.query(models.Paquete).filter_by(foto_url=url).first()
    announcement = db.query(models.Comunicado).filter_by(adjunto_url=url).first()
    allowed = False
    if incident:
        allowed = user.es_super_admin or (is_admin(user) and can_access_set(db, user, incident.conjunto_id)) or incident.solicitante_id == user.id
    elif package:
        unit = db.query(models.Inmueble).filter_by(id=package.inmueble_id).first()
        allowed = bool(unit and (user.es_super_admin or can_access_set(db, user, unit.conjunto_id) or user.id in (unit.propietario_id, unit.inquilino_id)))
    elif announcement:
        allowed = can_access_set(db, user, announcement.conjunto_id)
    if not allowed: raise HTTPException(404, "Archivo no encontrado")
    root = Path(os.getenv("UPLOAD_DIR", "./uploads")).resolve(); path = (root / filename).resolve()
    if root not in path.parents or not path.is_file(): raise HTTPException(404, "Archivo no encontrado")
    from fastapi.responses import FileResponse
    return FileResponse(path, headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"})
