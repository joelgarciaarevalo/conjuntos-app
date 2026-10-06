from datetime import datetime, date
import os
from zoneinfo import ZoneInfo
import csv
import io
from decimal import Decimal
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy.orm import Session
from app.database import get_db
from app import models
from app.dependencies import get_usuario_actual, hashear_password
from app.permissions import allowed_set_ids, require_admin_set, require_finance_set, require_staff_set
from app.security_tokens import issue_token
from app.services.notifications import deliver_external

router = APIRouter(tags=["Gestión del conjunto"])

def now_local():
    return datetime.now(ZoneInfo("America/Bogota")).replace(tzinfo=None)


def admin(u: models.Usuario):
    return bool(u.es_super_admin or u.rol == "admin")


def require_admin(u: models.Usuario, db=None, conjunto_id=None):
    if not admin(u):
        raise HTTPException(403, "Solo administración puede realizar esta acción")
    if conjunto_id is not None:
        require_admin_set(db, u, conjunto_id)


def add_notification(db, user_id: int, title: str, message: str):
    db.add(models.Notificacion(usuario_id=user_id, titulo=title, mensaje=message))
    from app.services.notifications import deliver_external
    deliver_external(db, user_id, title, message)


@router.get("/notificaciones/")
def notifications(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    return db.query(models.Notificacion).filter(models.Notificacion.usuario_id == user.id).order_by(models.Notificacion.id.desc()).limit(100).all()


@router.patch("/notificaciones/{notification_id}/leida")
def read_notification(notification_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Notificacion).filter_by(id=notification_id, usuario_id=user.id).first()
    if not row:
        raise HTTPException(404, "Notificación no encontrada")
    row.leida = True
    db.commit()
    return {"ok": True}


@router.get("/gestion/usuarios")
def users(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user)
    if user.es_super_admin:
        return db.query(models.Usuario).order_by(models.Usuario.nombre_completo).all()
    ids = allowed_set_ids(db, user)
    unit_user_ids = {uid for row in db.query(models.Inmueble.propietario_id, models.Inmueble.inquilino_id).filter(models.Inmueble.conjunto_id.in_(ids)).all() for uid in row if uid}
    staff_ids = {x.usuario_id for x in db.query(models.ConjuntoMembership).filter(models.ConjuntoMembership.conjunto_id.in_(ids)).all()}
    visible = unit_user_ids | staff_ids
    return db.query(models.Usuario).filter(models.Usuario.id.in_(visible)).order_by(models.Usuario.nombre_completo).all() if visible else []


@router.patch("/gestion/usuarios/{user_id}/rol")
def set_user_role(user_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user)
    role = data.get("rol")
    if role not in ("residente", "portero", "admin"):
        raise HTTPException(422, "Rol inválido")
    target = db.query(models.Usuario).filter_by(id=user_id).first()
    if not target:
        raise HTTPException(404, "Usuario no encontrado")
    if target.es_super_admin and not user.es_super_admin:
        raise HTTPException(403, "No puedes cambiar el rol del super administrador")
    if not user.es_super_admin:
        if role != "residente": raise HTTPException(403, "Solo superadministración puede asignar roles administrativos o de portería")
        ids = allowed_set_ids(db, user)
        linked = db.query(models.Inmueble.id).filter(models.Inmueble.conjunto_id.in_(ids)).filter((models.Inmueble.propietario_id == target.id) | (models.Inmueble.inquilino_id == target.id)).first()
        if not linked: raise HTTPException(403, "El usuario no pertenece a tus conjuntos")
    target.rol = role
    db.commit()
    return {"id": target.id, "rol": target.rol}


@router.post("/gestion/conjuntos/{conjunto_id}/contadores", status_code=201)
def crear_cuenta_contador(conjunto_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    """Provisiona un usuario de contador independiente y restringido a un conjunto."""
    require_admin(user, db, conjunto_id)
    nombre = str(data.get("nombre_completo") or "").strip()
    cedula = str(data.get("cedula") or "").strip()
    email = str(data.get("email") or "").strip().lower()
    password = str(data.get("password") or "")
    telefono = str(data.get("telefono") or "").strip() or None
    if len(nombre) < 3 or len(nombre) > 100:
        raise HTTPException(422, "El nombre debe tener entre 3 y 100 caracteres")
    if not cedula or len(cedula) > 20:
        raise HTTPException(422, "La identificación del contador es obligatoria")
    if "@" not in email or len(email) > 100:
        raise HTTPException(422, "Ingresa un correo válido")
    if len(password) < 12:
        raise HTTPException(422, "La contraseña temporal debe tener al menos 12 caracteres")
    if db.query(models.Usuario.id).filter(models.Usuario.email == email).first():
        raise HTTPException(409, "Ya existe una cuenta con ese correo")
    if db.query(models.Usuario.id).filter(models.Usuario.cedula == cedula).first():
        raise HTTPException(409, "Ya existe una cuenta con esa identificación")
    production = os.getenv("APP_ENV", "development").lower() == "production"
    row = models.Usuario(nombre_completo=nombre, cedula=cedula, email=email,
        password_hash=hashear_password(password), telefono=telefono, rol="contador",
        email_verificado=not production, cambio_password_requerido=True)
    db.add(row); db.flush()
    db.add(models.ConjuntoMembership(usuario_id=row.id, conjunto_id=conjunto_id, rol="contador"))
    if production:
        token = issue_token(db, row, "verify", 24 * 60)
    db.commit()
    if production:
        link = f"{os.getenv('API_PUBLIC_URL', 'http://localhost:8000')}/auth/verify-email?token={token}"
        deliver_external(db, row.id, "Verifica tu cuenta de contador IRIS", f"Abre este enlace para activar tu acceso: {link}", allow_unverified=True)
    return {"id": row.id, "nombre_completo": row.nombre_completo, "email": row.email,
        "rol": "contador", "conjunto_id": conjunto_id, "cambio_password_requerido": True,
        "verificacion_requerida": production}


@router.post("/conjuntos/{conjunto_id}/inmuebles", status_code=201)
def create_unit(conjunto_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user, db, conjunto_id)
    if not db.query(models.Conjunto).filter_by(id=conjunto_id).first():
        raise HTTPException(404, "Conjunto no encontrado")
    owner = db.query(models.Usuario).filter_by(id=data.get("propietario_id")).first()
    if not owner:
        raise HTTPException(404, "Propietario no encontrado")
    estado = str(data.get("estado", "ocupado")).lower()
    if estado not in {e.value for e in models.EstadoInmueble}:
        raise HTTPException(422, "Estado de inmueble inválido")
    unit = models.Inmueble(conjunto_id=conjunto_id, torre=data.get("torre"), numero=data.get("numero", "").strip(), propietario_id=owner.id, inquilino_id=data.get("inquilino_id"), estado=estado)
    if not unit.numero:
        raise HTTPException(422, "El número de inmueble es obligatorio")
    db.add(unit); db.commit(); db.refresh(unit)
    return unit


@router.patch("/conjuntos/{conjunto_id}/inmuebles/{unit_id}")
def update_unit(conjunto_id: int, unit_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user, db, conjunto_id)
    row = db.query(models.Inmueble).filter_by(id=unit_id, conjunto_id=conjunto_id).first()
    if not row: raise HTTPException(404, "Inmueble no encontrado")
    for key in ("torre", "numero", "propietario_id", "inquilino_id", "estado"):
        if key in data: setattr(row, key, str(data[key]).lower() if key == "estado" else data[key])
    if row.estado not in {e.value for e in models.EstadoInmueble}: raise HTTPException(422, "Estado inválido")
    db.commit(); db.refresh(row); return row


@router.delete("/conjuntos/{conjunto_id}/inmuebles/{unit_id}", status_code=204)
def delete_unit(conjunto_id: int, unit_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user, db, conjunto_id)
    row = db.query(models.Inmueble).filter_by(id=unit_id, conjunto_id=conjunto_id).first()
    if not row: raise HTTPException(404, "Inmueble no encontrado")
    if db.query(models.Reserva).filter_by(inmueble_id=unit_id).first() or db.query(models.PQRS).filter_by(inmueble_id=unit_id).first():
        raise HTTPException(409, "El inmueble tiene reservas o PQRS asociadas; no se puede eliminar")
    db.delete(row); db.commit(); return Response(status_code=204)


@router.post("/conjuntos/{conjunto_id}/zonas-comunes", status_code=201)
def create_zone(conjunto_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user, db, conjunto_id)
    if not db.query(models.Conjunto).filter_by(id=conjunto_id).first(): raise HTTPException(404, "Conjunto no encontrado")
    name = str(data.get("nombre", "")).strip()
    opening = str(data.get("hora_apertura", "06:00")); closing = str(data.get("hora_cierre", "22:00"))
    if not name: raise HTTPException(422, "El nombre es obligatorio")
    try:
        datetime.strptime(opening, "%H:%M"); datetime.strptime(closing, "%H:%M")
    except ValueError: raise HTTPException(422, "Horario debe usar HH:MM")
    if opening >= closing: raise HTTPException(422, "La hora de cierre debe ser posterior a apertura")
    duration = int(data.get("duracion_max_minutos", 240))
    if not 1 <= duration <= 1440: raise HTTPException(422, "Duración máxima debe ser de 1 a 1440 minutos")
    row = models.ZonaComun(conjunto_id=conjunto_id, nombre=name, requiere_aprobacion=bool(data.get("requiere_aprobacion", True)), hora_apertura=opening, hora_cierre=closing, duracion_max_minutos=duration, max_dias_anticipacion=int(data.get("max_dias_anticipacion", 30)), max_reservas_dia_inmueble=int(data.get("max_reservas_dia_inmueble", 2)))
    db.add(row); db.commit(); db.refresh(row); return row


@router.patch("/conjuntos/{conjunto_id}/zonas-comunes/{zone_id}")
def update_zone(conjunto_id: int, zone_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user, db, conjunto_id)
    row = db.query(models.ZonaComun).filter_by(id=zone_id, conjunto_id=conjunto_id).first()
    if not row: raise HTTPException(404, "Zona no encontrada")
    for key in ("nombre", "requiere_aprobacion", "hora_apertura", "hora_cierre", "duracion_max_minutos", "max_dias_anticipacion", "max_reservas_dia_inmueble"):
        if key in data: setattr(row, key, data[key])
    if row.hora_apertura >= row.hora_cierre or not 1 <= row.duracion_max_minutos <= 1440 or not 1 <= row.max_dias_anticipacion <= 365 or not 1 <= row.max_reservas_dia_inmueble <= 20: raise HTTPException(422, "Horario o duración inválidos")
    db.commit(); db.refresh(row); return row


@router.delete("/conjuntos/{conjunto_id}/zonas-comunes/{zone_id}", status_code=204)
def delete_zone(conjunto_id: int, zone_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user, db, conjunto_id)
    row = db.query(models.ZonaComun).filter_by(id=zone_id, conjunto_id=conjunto_id).first()
    if not row: raise HTTPException(404, "Zona no encontrada")
    if db.query(models.Reserva).filter_by(zona_id=zone_id).first(): raise HTTPException(409, "La zona tiene reservas asociadas")
    db.delete(row); db.commit(); return Response(status_code=204)


@router.get("/comunicados/")
def list_announcements(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    units = db.query(models.Inmueble).filter((models.Inmueble.propietario_id == user.id) | (models.Inmueble.inquilino_id == user.id)).all()
    ids = {i.conjunto_id for i in units}
    q = db.query(models.Comunicado).filter(models.Comunicado.publicado.is_(True))
    if admin(user): q = q.filter(models.Comunicado.conjunto_id.in_(allowed_set_ids(db, user)))
    else: q = q.filter(models.Comunicado.conjunto_id.in_(ids))
    return q.order_by(models.Comunicado.id.desc()).all()


@router.post("/comunicados/", status_code=201)
def create_announcement(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    cid = int(data.get("conjunto_id", 0)); require_admin(user, db, cid)
    title = str(data.get("titulo", "")).strip(); content = str(data.get("contenido", "")).strip()
    if not title or not content or not db.query(models.Conjunto).filter_by(id=cid).first(): raise HTTPException(422, "Indica conjunto, título y contenido")
    row = models.Comunicado(conjunto_id=cid, titulo=title, contenido=content, creado_por=user.id)
    db.add(row)
    units = db.query(models.Inmueble).filter_by(conjunto_id=cid).all()
    resident_ids = {x for i in units for x in (i.propietario_id, i.inquilino_id) if x}
    for uid in resident_ids: add_notification(db, uid, "Nuevo comunicado", title)
    db.commit(); db.refresh(row); return row


@router.get("/cuotas/")
def list_fees(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    q = db.query(models.Cuota)
    if admin(user) or user.rol == "contador":
        ids = allowed_set_ids(db, user); unit_ids = [x.id for x in db.query(models.Inmueble.id).filter(models.Inmueble.conjunto_id.in_(ids)).all()]
        q = q.filter(models.Cuota.inmueble_id.in_(unit_ids))
    else:
        unit_ids = [i.id for i in db.query(models.Inmueble).filter((models.Inmueble.propietario_id == user.id) | (models.Inmueble.inquilino_id == user.id)).all()]
        q = q.filter(models.Cuota.inmueble_id.in_(unit_ids))
    return q.order_by(models.Cuota.fecha_vencimiento.desc()).all()


@router.post("/cuotas/", status_code=201)
def create_fee(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    try:
        amount = Decimal(str(data.get("monto"))); due = date.fromisoformat(str(data.get("fecha_vencimiento")))
        unit_id = int(data.get("inmueble_id")); period = str(data.get("periodo", "")); title = str(data.get("concepto", "Cuota de administración"))
    except Exception: raise HTTPException(422, "Datos de cuota inválidos")
    unit = db.query(models.Inmueble).filter_by(id=unit_id).first()
    if not unit or amount <= 0: raise HTTPException(422, "Inmueble o monto inválido")
    if user.rol == "contador": require_finance_set(db, user, unit.conjunto_id)
    else: require_admin(user, db, unit.conjunto_id)
    row = models.Cuota(inmueble_id=unit_id, concepto=title, periodo=period, monto=amount, fecha_vencimiento=due, creado_por=user.id)
    db.add(row)
    for uid in {unit.propietario_id, unit.inquilino_id} - {None}: add_notification(db, uid, "Nueva cuota", f"{title} · {period} · ${amount}")
    db.commit(); db.refresh(row); return row


@router.patch("/cuotas/{fee_id}/pago")
def update_fee(fee_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Cuota).filter_by(id=fee_id).first()
    if not row: raise HTTPException(404, "Cuota no encontrada")
    unit = db.query(models.Inmueble).filter_by(id=row.inmueble_id).first()
    if user.rol == "contador": require_finance_set(db, user, unit.conjunto_id)
    else: require_admin(user, db, unit.conjunto_id)
    paid = bool(data.get("pagada", True)); row.estado = "pagada" if paid else "pendiente"; row.pagada_en = now_local() if paid else None
    db.commit(); db.refresh(row); return row


@router.get("/cuotas/paz-y-salvo/{inmueble_id}")
def clearance(inmueble_id: int, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    unit = db.query(models.Inmueble).filter_by(id=inmueble_id).first()
    if not unit: raise HTTPException(404, "Inmueble no encontrado")
    if admin(user): require_admin(user, db, unit.conjunto_id)
    elif user.rol == "contador": require_finance_set(db, user, unit.conjunto_id)
    elif user.id not in (unit.propietario_id, unit.inquilino_id): raise HTTPException(403, "Sin acceso")
    pending = db.query(models.Cuota).filter_by(inmueble_id=inmueble_id, estado="pendiente").all()
    return {"inmueble_id": inmueble_id, "paz_y_salvo": len(pending) == 0, "cuotas_pendientes": len(pending), "saldo": str(sum((Decimal(str(x.monto)) for x in pending), Decimal("0")))}


@router.get("/paquetes/")
def packages(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    q = db.query(models.Paquete)
    if admin(user) or user.rol == "portero":
        ids = allowed_set_ids(db, user); q = q.join(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids))
    else:
        ids = [i.id for i in db.query(models.Inmueble).filter((models.Inmueble.propietario_id == user.id) | (models.Inmueble.inquilino_id == user.id)).all()]
        q = q.filter(models.Paquete.inmueble_id.in_(ids))
    return q.order_by(models.Paquete.id.desc()).all()


@router.post("/paquetes/", status_code=201)
def receive_package(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    if user.rol not in ("portero", "admin") and not user.es_super_admin: raise HTTPException(403, "Solo portería o administración registra paquetes")
    unit = db.query(models.Inmueble).filter_by(id=data.get("inmueble_id")).first()
    if not unit: raise HTTPException(404, "Inmueble no encontrado")
    require_staff_set(db, user, unit.conjunto_id)
    row = models.Paquete(inmueble_id=unit.id, transportadora=data.get("transportadora"), descripcion=str(data.get("descripcion", "Paquete")), recibido_por=user.id)
    db.add(row)
    for uid in {unit.propietario_id, unit.inquilino_id} - {None}: add_notification(db, uid, "Paquete recibido", f"{row.descripcion} · {row.transportadora or 'sin transportadora'}")
    db.commit(); db.refresh(row); return row


@router.patch("/paquetes/{package_id}/retirar")
def pickup_package(package_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    row = db.query(models.Paquete).filter_by(id=package_id).first()
    if not row: raise HTTPException(404, "Paquete no encontrado")
    unit = db.query(models.Inmueble).filter_by(id=row.inmueble_id).first()
    if user.rol in ("portero", "admin") or user.es_super_admin: require_staff_set(db, user, unit.conjunto_id)
    if user.rol not in ("portero", "admin") and not user.es_super_admin and user.id not in (unit.propietario_id, unit.inquilino_id): raise HTTPException(403, "Sin permiso para retirar este paquete")
    if row.retirado: raise HTTPException(409, "El paquete ya fue retirado")
    row.retirado = True; row.retirado_por = str(data.get("retirado_por") or user.nombre_completo); row.retirado_en = now_local()
    db.commit(); return {"ok": True}


@router.get("/parqueaderos/")
def parking(db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    q = db.query(models.Parqueadero)
    if admin(user) or user.rol == "portero": q = q.filter(models.Parqueadero.conjunto_id.in_(allowed_set_ids(db, user)))
    else:
        ids = [i.id for i in db.query(models.Inmueble).filter((models.Inmueble.propietario_id == user.id) | (models.Inmueble.inquilino_id == user.id)).all()]
        q = q.filter(models.Parqueadero.inmueble_id.in_(ids), models.Parqueadero.activo.is_(True))
    return q.order_by(models.Parqueadero.identificador).all()


@router.post("/parqueaderos/", status_code=201)
def create_parking(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    cid = int(data.get("conjunto_id", 0)); require_admin(user, db, cid)
    unit_id = data.get("inmueble_id")
    if unit_id and not db.query(models.Inmueble.id).filter_by(id=unit_id, conjunto_id=cid).first(): raise HTTPException(422, "El inmueble no pertenece al conjunto")
    identificador = str(data.get("identificador", "")).strip()
    if db.query(models.Parqueadero.id).filter_by(conjunto_id=cid, identificador=identificador).first(): raise HTTPException(409, "Ese identificador de parqueadero ya existe en el conjunto")
    try: tarifa = Decimal(str(data.get("tarifa_mensual") or 0))
    except Exception: raise HTTPException(422, "La tarifa mensual debe ser un número válido")
    if tarifa < 0: raise HTTPException(422, "La tarifa mensual no puede ser negativa")
    row = models.Parqueadero(conjunto_id=cid, inmueble_id=unit_id, identificador=identificador, tipo=str(data.get("tipo", "privado")), placa=data.get("placa"), tarifa_mensual=tarifa)
    if not row.identificador: raise HTTPException(422, "Identificador requerido")
    db.add(row); db.commit(); db.refresh(row); return row


@router.post("/parqueaderos/generar-cobros")
def generate_parking_charges(data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    cid = int(data.get("conjunto_id") or 0)
    if user.rol == "contador": require_finance_set(db, user, cid)
    else: require_admin(user, db, cid)
    periodo = str(data.get("periodo") or "").strip()
    try:
        datetime.strptime(periodo, "%Y-%m")
        vencimiento = date.fromisoformat(str(data.get("fecha_vencimiento") or ""))
    except ValueError:
        raise HTTPException(422, "Indica el periodo YYYY-MM y una fecha de vencimiento válida")
    filas = db.query(models.Parqueadero).filter(
        models.Parqueadero.conjunto_id == cid,
        models.Parqueadero.activo.is_(True),
        models.Parqueadero.inmueble_id.is_not(None),
        models.Parqueadero.tarifa_mensual > 0,
    ).order_by(models.Parqueadero.id).all()
    creados = omitidos = 0
    for parking in filas:
        concepto = f"Parqueadero {parking.identificador}"
        existe = db.query(models.Cuota.id).filter_by(inmueble_id=parking.inmueble_id, concepto=concepto, periodo=periodo).first()
        if existe:
            omitidos += 1
            continue
        cuota = models.Cuota(
            inmueble_id=parking.inmueble_id, concepto=concepto, periodo=periodo,
            monto=parking.tarifa_mensual, fecha_vencimiento=vencimiento,
            estado="pendiente", creado_por=user.id,
        )
        db.add(cuota)
        inmueble = db.query(models.Inmueble).filter_by(id=parking.inmueble_id).first()
        if inmueble:
            for uid in {inmueble.propietario_id, inmueble.inquilino_id} - {None}:
                add_notification(db, uid, "Nueva cuota de parqueadero", f"Se generó el cobro de {concepto} para el periodo {periodo}.")
        creados += 1
    db.commit()
    return {"conjunto_id": cid, "periodo": periodo, "creados": creados, "omitidos": omitidos}


@router.patch("/parqueaderos/{parking_id}")
def update_parking(parking_id: int, data: dict, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user)
    row = db.query(models.Parqueadero).filter_by(id=parking_id).first()
    if not row: raise HTTPException(404, "Parqueadero no encontrado")
    require_admin(user, db, row.conjunto_id)
    if "inmueble_id" in data and data["inmueble_id"] is not None and not db.query(models.Inmueble.id).filter_by(id=data["inmueble_id"], conjunto_id=row.conjunto_id).first(): raise HTTPException(422, "El inmueble no pertenece al conjunto")
    if "tarifa_mensual" in data:
        try: tarifa = Decimal(str(data["tarifa_mensual"]))
        except Exception: raise HTTPException(422, "La tarifa mensual debe ser un número válido")
        if tarifa < 0: raise HTTPException(422, "La tarifa mensual no puede ser negativa")
        row.tarifa_mensual = tarifa
    for key in ("inmueble_id", "identificador", "tipo", "placa", "activo"):
        if key in data: setattr(row, key, data[key])
    db.commit(); db.refresh(row); return row


@router.get("/reportes/{report_type}.csv")
def report(report_type: str, desde: Optional[str] = None, hasta: Optional[str] = None, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user)
    ids = allowed_set_ids(db, user)
    mapping = {"reservas": models.Reserva, "pqrs": models.PQRS, "accesos": None}
    if report_type not in mapping: raise HTTPException(404, "Reporte no disponible")
    out = io.StringIO(); writer = csv.writer(out); writer.writerow(["tipo", "id", "fecha", "detalle", "estado"])
    start = datetime.fromisoformat(desde) if desde else None; end = datetime.fromisoformat(hasta) if hasta else None
    if report_type == "reservas":
        for x in db.query(models.Reserva).join(models.ZonaComun).filter(models.ZonaComun.conjunto_id.in_(ids)).order_by(models.Reserva.fecha_inicio).all():
            if start and x.fecha_inicio < start or end and x.fecha_inicio > end: continue
            writer.writerow(["reserva", x.id, x.fecha_inicio.isoformat(), f"zona={x.zona_id}; inmueble={x.inmueble_id}; fin={x.fecha_fin.isoformat()}", getattr(x.estado, "value", x.estado)])
    elif report_type == "pqrs":
        for x in db.query(models.PQRS).join(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids)).order_by(models.PQRS.fecha_radicacion).all():
            if start and x.fecha_radicacion < start or end and x.fecha_radicacion > end: continue
            writer.writerow(["pqrs", x.id, x.fecha_radicacion.isoformat(), f"tipo={getattr(x.tipo,'value',x.tipo)}; {x.descripcion}; respuesta={x.respuesta_admin or ''}", getattr(x.estado, "value", x.estado)])
    else:
        for x in db.query(__import__("app.models_qr", fromlist=["AccesoQR"]).AccesoQR).join(__import__("app.models_qr", fromlist=["CodigoQR"]).CodigoQR, __import__("app.models_qr", fromlist=["AccesoQR"]).AccesoQR.codigo_id == __import__("app.models_qr", fromlist=["CodigoQR"]).CodigoQR.id).join(models.Inmueble, __import__("app.models_qr", fromlist=["CodigoQR"]).CodigoQR.inmueble_id == models.Inmueble.id).filter(models.Inmueble.conjunto_id.in_(ids)).order_by(__import__("app.models_qr", fromlist=["AccesoQR"]).AccesoQR.fecha).all():
            if start and x.fecha < start or end and x.fecha > end: continue
            writer.writerow(["acceso", x.id, x.fecha.isoformat(), x.motivo, x.resultado])
    return Response(content="\ufeff" + out.getvalue(), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f"attachment; filename=iris-{report_type}.csv"})


@router.get("/reportes/{report_type}.xlsx")
def report_xlsx(report_type: str, db: Session = Depends(get_db), user=Depends(get_usuario_actual)):
    require_admin(user)
    ids = allowed_set_ids(db, user)
    if report_type not in ("reservas", "pqrs", "accesos"):
        raise HTTPException(404, "Reporte no disponible")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from app.models_qr import AccesoQR
    book = Workbook(); sheet = book.active; sheet.title = report_type.capitalize()
    sheet.append(["Tipo", "ID", "Fecha", "Detalle", "Estado"])
    if report_type == "reservas":
        for x in db.query(models.Reserva).join(models.ZonaComun).filter(models.ZonaComun.conjunto_id.in_(ids)).order_by(models.Reserva.fecha_inicio).all():
            sheet.append(["Reserva", x.id, x.fecha_inicio, f"Zona {x.zona_id}; inmueble {x.inmueble_id}; fin {x.fecha_fin}", getattr(x.estado, "value", x.estado)])
    elif report_type == "pqrs":
        for x in db.query(models.PQRS).join(models.Inmueble).filter(models.Inmueble.conjunto_id.in_(ids)).order_by(models.PQRS.fecha_radicacion).all():
            sheet.append(["PQRS", x.id, x.fecha_radicacion, f"{getattr(x.tipo, 'value', x.tipo)}: {x.descripcion}; respuesta: {x.respuesta_admin or ''}", getattr(x.estado, "value", x.estado)])
    else:
        for x in db.query(AccesoQR).join(__import__("app.models_qr", fromlist=["CodigoQR"]).CodigoQR, AccesoQR.codigo_id == __import__("app.models_qr", fromlist=["CodigoQR"]).CodigoQR.id).join(models.Inmueble, __import__("app.models_qr", fromlist=["CodigoQR"]).CodigoQR.inmueble_id == models.Inmueble.id).filter(models.Inmueble.conjunto_id.in_(ids)).order_by(AccesoQR.fecha).all():
            sheet.append(["Acceso", x.id, x.fecha, x.motivo, x.resultado])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0C3B45")
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        width = min(60, max(14, max(len(str(c.value or "")) for c in col) + 2))
        sheet.column_dimensions[get_column_letter(col[0].column)].width = width
    output = io.BytesIO(); book.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=iris-{report_type}.xlsx"})
